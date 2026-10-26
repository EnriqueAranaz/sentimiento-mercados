"""Carga semilla de la tabla `empresas`.

Universo (en_universo = true): Ibex 35, DAX 40, Euro Stoxx 50, Stoxx Europe 50 y ~50 mayores de EE. UU.
Más todas las empresas de la hoja `afectaciones` de referencia/clasificacion_titulares_2026-10-26_v3.xlsx,
con su sector, ticker y en_universo TAL COMO ESTÁN en el Excel (si una empresa está en ambos, manda el Excel).

La composición de los índices es aproximada (octubre de 2026): revisar en referencia/empresas_cargadas.xlsx.

Uso:
    python clasificacion/semilla_empresas.py            # carga en Supabase (upsert por nombre) y exporta el Excel
    python clasificacion/semilla_empresas.py --solo-excel   # solo genera el Excel de revisión
"""
from __future__ import annotations

import os
import sys
from pathlib import Path

from dotenv import load_dotenv
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

RAIZ = Path(__file__).resolve().parent.parent
REFERENCIA = RAIZ / "referencia" / "clasificacion_titulares_2026-10-26_v3.xlsx"
SALIDA = RAIZ / "referencia" / "empresas_cargadas.xlsx"

# (nombre, ticker TradingView, sector). El mercado se deduce del índice de la lista.
IBEX = [
    ("Acciona Energía", "BME:ANE", "Eléctricas, renovables y redes"),
    ("Aena", "BME:AENA", "Transporte y turismo"),
    ("ArcelorMittal", "BME:MTS", "Materiales, química y minería"),
    ("Cellnex", "BME:CLNX", "Telecomunicaciones"),
    ("Endesa", "BME:ELE", "Eléctricas, renovables y redes"),
    ("Fluidra", "BME:FDR", "Industria, defensa y aeroespacial"),
    ("Grifols", "BME:GRF", "Salud y farmacia"),
    ("Logista", "BME:LOG", "Consumo y distribución"),
    ("Mapfre", "BME:MAP", "Seguros y servicios financieros"),
    ("Sacyr", "BME:SCYR", "Construcción e infraestructuras"),
    # Las otras 25 del Ibex ya están en el Excel con en_universo = Sí.
]
DAX = [
    ("Adidas", "XETR:ADS", "Lujo, moda y cosmética"),
    ("Allianz", "XETR:ALV", "Seguros y servicios financieros"),
    ("BMW", "XETR:BMW", "Automóvil"),
    ("Beiersdorf", "XETR:BEI", "Lujo, moda y cosmética"),
    ("Brenntag", "XETR:BNR", "Materiales, química y minería"),
    ("Commerzbank", "XETR:CBK", "Banca"),
    ("Continental", "XETR:CON", "Automóvil"),
    ("Daimler Truck", "XETR:DTG", "Automóvil"),
    ("Deutsche Bank", "XETR:DBK", "Banca"),
    ("Deutsche Börse", "XETR:DB1", "Seguros y servicios financieros"),
    ("DHL Group", "XETR:DHL", "Transporte y turismo"),
    ("E.ON", "XETR:EOAN", "Eléctricas, renovables y redes"),
    ("Fresenius", "XETR:FRE", "Salud y farmacia"),
    ("Fresenius Medical Care", "XETR:FME", "Salud y farmacia"),
    ("GEA Group", "XETR:G1A", "Industria, defensa y aeroespacial"),
    ("Hannover Rück", "XETR:HNR1", "Seguros y servicios financieros"),
    ("Heidelberg Materials", "XETR:HEI", "Materiales, química y minería"),
    ("Henkel", "XETR:HEN3", "Consumo y distribución"),
    ("Infineon", "XETR:IFX", "Semiconductores"),
    ("Mercedes-Benz", "XETR:MBG", "Automóvil"),
    ("Merck KGaA", "XETR:MRK", "Salud y farmacia"),
    ("MTU Aero Engines", "XETR:MTX", "Industria, defensa y aeroespacial"),
    ("Munich Re", "XETR:MUV2", "Seguros y servicios financieros"),
    ("Porsche AG", "XETR:P911", "Automóvil"),
    ("Porsche SE", "XETR:PAH3", "Automóvil"),
    ("Qiagen", "XETR:QIA", "Salud y farmacia"),
    ("Rheinmetall", "XETR:RHM", "Industria, defensa y aeroespacial"),
    ("RWE", "XETR:RWE", "Eléctricas, renovables y redes"),
    ("Scout24", "XETR:G24", "Tecnología"),
    ("Siemens", "XETR:SIE", "Industria, defensa y aeroespacial"),
    ("Siemens Energy", "XETR:ENR", "Industria, defensa y aeroespacial"),
    ("Siemens Healthineers", "XETR:SHL", "Salud y farmacia"),
    ("Symrise", "XETR:SY1", "Materiales, química y minería"),
    ("Zalando", "XETR:ZAL", "Consumo y distribución"),
    # BASF, Bayer, Deutsche Telekom, SAP, Vonovia y Volkswagen ya están en el Excel. Airbus va en Europa.
]
EUROPA = [  # Euro Stoxx 50 y Stoxx Europe 50 (sin las ya listadas ni las del Excel)
    ("Airbus", "EURONEXT:AIR", "Industria, defensa y aeroespacial"),
    ("TotalEnergies", "EURONEXT:TTE", "Petróleo y gas"),
    ("Schneider Electric", "EURONEXT:SU", "Industria, defensa y aeroespacial"),
    ("Sanofi", "EURONEXT:SAN", "Salud y farmacia"),
    ("Air Liquide", "EURONEXT:AI", "Materiales, química y minería"),
    ("BNP Paribas", "EURONEXT:BNP", "Banca"),
    ("AXA", "EURONEXT:CS", "Seguros y servicios financieros"),
    ("Vinci", "EURONEXT:DG", "Construcción e infraestructuras"),
    ("Pernod Ricard", "EURONEXT:RI", "Consumo y distribución"),
    ("Saint-Gobain", "EURONEXT:SGO", "Construcción e infraestructuras"),
    ("Prosus", "EURONEXT:PRX", "Tecnología"),
    ("ING", "EURONEXT:INGA", "Banca"),
    ("Adyen", "EURONEXT:ADYEN", "Seguros y servicios financieros"),
    ("Ahold Delhaize", "EURONEXT:AD", "Consumo y distribución"),
    ("AB InBev", "EURONEXT:ABI", "Consumo y distribución"),
    ("UniCredit", "MIL:UCG", "Banca"),
    ("Intesa Sanpaolo", "MIL:ISP", "Banca"),
    ("Generali", "MIL:G", "Seguros y servicios financieros"),
    ("Ferrari", "MIL:RACE", "Automóvil"),
    ("Enel", "MIL:ENEL", "Eléctricas, renovables y redes"),
    ("Stellantis", "MIL:STLAM", "Automóvil"),
    ("Nordea", "OMXHEX:NDA_FI", "Banca"),
    ("Shell", "LSE:SHEL", "Petróleo y gas"),
    ("Unilever", "LSE:ULVR", "Consumo y distribución"),
    ("British American Tobacco", "LSE:BATS", "Consumo y distribución"),
    ("Diageo", "LSE:DGE", "Consumo y distribución"),
    ("RELX", "LSE:REL", "Medios y entretenimiento"),
    ("London Stock Exchange Group", "LSE:LSEG", "Seguros y servicios financieros"),
    ("Lloyds Banking Group", "LSE:LLOY", "Banca"),
    ("Rolls-Royce", "LSE:RR.", "Industria, defensa y aeroespacial"),
    ("National Grid", "LSE:NG.", "Eléctricas, renovables y redes"),
    ("Reckitt", "LSE:RKT", "Consumo y distribución"),
    ("Compass Group", "LSE:CPG", "Consumo y distribución"),
    ("Zurich Insurance", "SIX:ZURN", "Seguros y servicios financieros"),
    ("ABB", "SIX:ABBN", "Industria, defensa y aeroespacial"),
    ("Richemont", "SIX:CFR", "Lujo, moda y cosmética"),
    ("Holcim", "SIX:HOLN", "Materiales, química y minería"),
    ("Sika", "SIX:SIKA", "Materiales, química y minería"),
    ("Lonza", "SIX:LONN", "Salud y farmacia"),
    ("Novo Nordisk", "OMXCOP:NOVO_B", "Salud y farmacia"),
    ("Equinor", "OSL:EQNR", "Petróleo y gas"),
    ("Atlas Copco", "OMXSTO:ATCO_A", "Industria, defensa y aeroespacial"),
    ("Investor AB", "OMXSTO:INVE_B", "Seguros y servicios financieros"),
]
EEUU = [  # ~50 mayores por capitalización (sin las del Excel)
    ("Broadcom", "NASDAQ:AVGO", "Semiconductores"),
    ("Berkshire Hathaway", "NYSE:BRK.B", "Seguros y servicios financieros"),
    ("Walmart", "NASDAQ:WMT", "Consumo y distribución"),
    ("Visa", "NYSE:V", "Seguros y servicios financieros"),
    ("Mastercard", "NYSE:MA", "Seguros y servicios financieros"),
    ("Exxon Mobil", "NYSE:XOM", "Petróleo y gas"),
    ("Costco", "NASDAQ:COST", "Consumo y distribución"),
    ("Home Depot", "NYSE:HD", "Consumo y distribución"),
    ("Palantir", "NASDAQ:PLTR", "Tecnología"),
    ("AbbVie", "NYSE:ABBV", "Salud y farmacia"),
    ("Bank of America", "NYSE:BAC", "Banca"),
    ("UnitedHealth", "NYSE:UNH", "Salud y farmacia"),
    ("Philip Morris", "NYSE:PM", "Consumo y distribución"),
    ("GE Aerospace", "NYSE:GE", "Industria, defensa y aeroespacial"),
    ("Cisco", "NASDAQ:CSCO", "Tecnología"),
    ("Salesforce", "NYSE:CRM", "Tecnología"),
    ("Wells Fargo", "NYSE:WFC", "Banca"),
    ("Morgan Stanley", "NYSE:MS", "Banca"),
    ("Goldman Sachs", "NYSE:GS", "Banca"),
    ("Citigroup", "NYSE:C", "Banca"),
    ("BlackRock", "NYSE:BLK", "Seguros y servicios financieros"),
    ("Abbott", "NYSE:ABT", "Salud y farmacia"),
    ("McDonald's", "NYSE:MCD", "Consumo y distribución"),
    ("Linde", "NASDAQ:LIN", "Materiales, química y minería"),
    ("Intuitive Surgical", "NASDAQ:ISRG", "Salud y farmacia"),
    ("Verizon", "NYSE:VZ", "Telecomunicaciones"),
    ("Thermo Fisher", "NYSE:TMO", "Salud y farmacia"),
    ("Caterpillar", "NYSE:CAT", "Industria, defensa y aeroespacial"),
    ("PepsiCo", "NASDAQ:PEP", "Consumo y distribución"),
    ("Walt Disney", "NYSE:DIS", "Medios y entretenimiento"),
    ("Uber", "NYSE:UBER", "Transporte y turismo"),
    ("ServiceNow", "NYSE:NOW", "Tecnología"),
    ("Qualcomm", "NASDAQ:QCOM", "Semiconductores"),
    ("Texas Instruments", "NASDAQ:TXN", "Semiconductores"),
    ("Adobe", "NASDAQ:ADBE", "Tecnología"),
    ("Booking Holdings", "NASDAQ:BKNG", "Transporte y turismo"),
    ("RTX", "NYSE:RTX", "Industria, defensa y aeroespacial"),
    ("Applied Materials", "NASDAQ:AMAT", "Semiconductores"),
    # Apple, Microsoft, Nvidia, Alphabet, Amazon, Meta, Tesla, JPMorgan, Eli Lilly, Oracle, Netflix, J&J, P&G,
    # Coca-Cola, Chevron, IBM, AMD, Merck, T-Mobile US, AT&T, American Express y Boeing ya están en el Excel.
]
LISTAS = [("Ibex 35", IBEX), ("DAX", DAX), ("Europa", EUROPA), ("EE. UU.", EEUU)]

# Propuestas del agente clasificador aprobadas por el usuario (26/10/2026): fuera del universo.
APROBADAS = [
    ("Elecnor", "BME:ENO", "Construcción e infraestructuras"),
    ("Coinbase", "NASDAQ:COIN", "Seguros y servicios financieros"),
    ("DoorDash", "NASDAQ:DASH", "Tecnología"),
    ("Geely", "HKEX:175", "Automóvil"),
    ("SAIC Motor", "SSE:600104", "Automóvil"),
    ("Chery", "HKEX:9973", "Automóvil"),  # verificado: cotiza en HKEX (9973) desde el 25/09/2025
    ("Woodside Energy", "ASX:WDS", "Petróleo y gas"),
    ("Hong Kong Exchanges and Clearing", "HKEX:388", "Seguros y servicios financieros"),
    ("Marriott Vacations Worldwide", "NYSE:VAC", "Transporte y turismo"),
    ("BeOne Medicines", "NASDAQ:ONC", "Salud y farmacia"),
    ("Sichuan Kelun-Biotech", "HKEX:6990", "Salud y farmacia"),
    ("Atico Mining", "TSXV:ATY", "Materiales, química y minería"),
    ("BlackRock Smaller Companies Trust", "LSE:BRSC", "Seguros y servicios financieros"),
]

# Alias habituales en la prensa española (además del propio nombre).
ALIAS_EXTRA = {
    "Santander": ["Banco Santander"],
    "Sabadell": ["Banco Sabadell"],
    "CaixaBank": ["Caixa Bank"],
    "Unicaja": ["Unicaja Banco"],
    "Telefónica": ["Telefonica"],
    "Redeia": ["Red Eléctrica", "REE"],
    "Merlin Properties": ["Merlin"],
    "Colonial": ["Inmobiliaria Colonial"],
    "IAG": ["International Airlines Group"],
    "Inditex": ["Zara"],
    "ArcelorMittal": ["Arcelor"],
    "Rovi": ["Laboratorios Rovi"],
    "Acciona Energía": ["Acciona Energia"],
    "Meta": ["Meta Platforms", "Facebook"],
    "Alphabet": ["Google"],
    "Amazon": ["Amazon.com"],
    "JPMorgan": ["JP Morgan", "JPMorgan Chase", "J.P. Morgan"],
    "Johnson & Johnson": ["J&J"],
    "Procter & Gamble": ["P&G"],
    "Coca-Cola": ["Coca Cola"],
    "Eli Lilly": ["Lilly"],
    "T-Mobile US": ["T-Mobile"],
    "Hermès": ["Hermes"],
    "L'Oréal": ["LOreal"],
    "Samsung": ["Samsung Electronics"],
    "SK Hynix": ["Hynix"],
    "Volkswagen": ["VW", "Grupo Volkswagen"],
    "Mercedes-Benz": ["Mercedes", "Mercedes Benz"],
    "Deutsche Börse": ["Deutsche Borse"],
    "DHL Group": ["DHL", "Deutsche Post"],
    "E.ON": ["EON"],
    "Hannover Rück": ["Hannover Re", "Hannover Rueck"],
    "Munich Re": ["Münchener Rück"],
    "Porsche AG": ["Porsche"],
    "MTU Aero Engines": ["MTU"],
    "TotalEnergies": ["Total Energies"],
    "BNP Paribas": ["BNP"],
    "Saint-Gobain": ["Saint Gobain"],
    "AB InBev": ["Anheuser-Busch InBev"],
    "Intesa Sanpaolo": ["Intesa"],
    "Stellantis": [],
    "Shell": [],
    "British American Tobacco": ["BAT"],
    "London Stock Exchange Group": ["LSEG", "Bolsa de Londres"],
    "Lloyds Banking Group": ["Lloyds"],
    "Rolls-Royce": ["Rolls Royce"],
    "Zurich Insurance": ["Zurich Seguros"],
    "Novo Nordisk": [],
    "Berkshire Hathaway": ["Berkshire"],
    "Exxon Mobil": ["ExxonMobil", "Exxon"],
    "Bank of America": ["BofA"],
    "Goldman Sachs": ["Goldman"],
    "Citigroup": ["Citi"],
    "McDonald's": ["McDonalds"],
    "Walt Disney": ["Disney"],
    "Booking Holdings": ["Booking"],
    "GE Aerospace": ["General Electric"],
    "Philip Morris": ["Philip Morris International"],
    "Thermo Fisher": ["Thermo Fisher Scientific"],
    "AB Volvo": ["Volvo"],
    "Nu Holdings": ["Nubank"],
    "SpaceX": ["Space X"],
    "Air France-KLM": ["Air France"],
    "Euroclear": [],
    "Geely": ["Geely Automobile"],
    "SAIC Motor": ["SAIC"],
    "Chery": ["Chery Automobile"],
    "Woodside Energy": ["Woodside"],
    "Hong Kong Exchanges and Clearing": ["HKEX", "Bolsa de Hong Kong"],
    "Marriott Vacations Worldwide": ["Marriott Vacations", "MVCI"],
    "BeOne Medicines": ["BeOne", "BeiGene"],
    "Sichuan Kelun-Biotech": ["Kelun-Biotech", "Kelun Biotech"],
}

# Notas para la revisión del Excel de salida.
NOTAS = {
    "Anthropic": "Ticker del Excel (NASDAQ:ANTHROPIC): comprobar si cotiza realmente.",
    "OpenAI": "Ticker del Excel (NASDAQ:OPENAI): comprobar si cotiza realmente.",
    "Merck": "Merck & Co. (EE. UU.). La alemana es «Merck KGaA».",
}


def mercado_por_ticker(ticker: str | None) -> str:
    bolsa = (ticker or "").split(":", 1)[0]
    if bolsa == "BME":
        return "Ibex 35"
    if bolsa == "XETR":
        return "DAX"
    if bolsa in {"EURONEXT", "MIL", "SIX", "LSE", "OMXHEX", "OMXSTO", "OMXCOP", "OSL"}:
        return "Europa"
    if bolsa in {"NASDAQ", "NYSE", "CBOE", "OTC"}:
        return "EE. UU."
    return "Otros"


def empresas_excel() -> dict[str, dict]:
    ws = load_workbook(REFERENCIA, data_only=True)["afectaciones"]
    filas = ws.iter_rows(values_only=True)
    cab = next(filas)
    out = {}
    for f in filas:
        d = dict(zip(cab, f))
        if d["ambito"] == "Empresa":
            out[d["objeto"]] = {
                "nombre": d["objeto"], "ticker": d["ticker"] or None, "sector": d["sector"],
                "en_universo": d["en_universo"] == "Sí", "mercado": mercado_por_ticker(d["ticker"]),
                "origen": "Excel de referencia",
            }
    return out


def construir() -> list[dict]:
    empresas = empresas_excel()
    for mercado, lista in LISTAS:
        for nombre, ticker, sector in lista:
            if nombre in empresas:
                empresas[nombre]["origen"] += f" + universo {mercado}"
                continue
            empresas[nombre] = {"nombre": nombre, "ticker": ticker, "sector": sector, "en_universo": True,
                                "mercado": mercado, "origen": f"Universo {mercado}"}
    for nombre, ticker, sector in APROBADAS:
        empresas.setdefault(nombre, {"nombre": nombre, "ticker": ticker, "sector": sector, "en_universo": False,
                                     "mercado": mercado_por_ticker(ticker),
                                     "origen": "Propuesta del agente, aprobada 26/10/2026"})
    for e in empresas.values():
        alias = [e["nombre"], *ALIAS_EXTRA.get(e["nombre"], [])]
        e["alias"] = list({a.casefold(): a for a in reversed(alias)}.values())[::-1]  # sin repetidos por mayúsculas
        e["nota"] = NOTAS.get(e["nombre"], "")
    return sorted(empresas.values(), key=lambda e: (not e["en_universo"], e["mercado"], e["nombre"]))


def exportar(empresas: list[dict]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "empresas"
    cab = ["nombre", "alias", "ticker", "sector", "mercado", "en_universo", "origen", "nota"]
    ws.append(cab)
    for c in ws[1]:
        c.font = Font(bold=True)
    for e in empresas:
        ws.append([e["nombre"], " | ".join(e["alias"]), e["ticker"], e["sector"], e["mercado"],
                   "Sí" if e["en_universo"] else "No", e["origen"], e["nota"]])
    for col, ancho in zip("ABCDEFGH", (28, 50, 18, 34, 10, 11, 34, 60)):
        ws.column_dimensions[col].width = ancho
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(SALIDA)


def main() -> int:
    empresas = construir()
    exportar(empresas)
    print(f"{len(empresas)} empresas · en universo: {sum(e['en_universo'] for e in empresas)} · Excel: {SALIDA}")
    if "--solo-excel" in sys.argv:
        return 0
    load_dotenv(RAIZ / ".env")
    from supabase import create_client

    db = create_client(os.environ["SUPABASE_URL"], os.environ["SUPABASE_SERVICE_ROLE_KEY"])
    filas = [{k: e[k] for k in ("nombre", "alias", "ticker", "sector", "mercado", "en_universo")} for e in empresas]
    for i in range(0, len(filas), 100):
        db.table("empresas").upsert(filas[i:i + 100], on_conflict="nombre").execute()
    total = db.table("empresas").select("id", count="exact").limit(1).execute().count
    print(f"Cargadas en Supabase. Filas en empresas: {total}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
