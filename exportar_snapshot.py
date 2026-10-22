"""Exporta una foto de la base de datos de Supabase a un Excel.

Genera exports/snapshot_AAAA-MM-DD_HHMM.xlsx (hora de Madrid) con:
  - "resumen": filas por tabla, artículos por fuente y día, y las últimas 10 ejecuciones.
  - una hoja por cada tabla del esquema public (se descubren solas vía la API de Supabase).

Las fechas se convierten a hora de Madrid y se guardan como fechas de Excel (dd/mm/aaaa hh:mm:ss).

Uso:
    python exportar_snapshot.py
Requiere: pip install -r requirements-export.txt  y  SUPABASE_URL / SUPABASE_SERVICE_ROLE_KEY en .env
"""
from __future__ import annotations

import os
import sys
from collections import Counter
from datetime import datetime
from pathlib import Path
from zoneinfo import ZoneInfo

import httpx
from dotenv import load_dotenv
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

RAIZ = Path(__file__).resolve().parent
MADRID = ZoneInfo("Europe/Madrid")
FORMATO_FECHA = "dd/mm/yyyy hh:mm:ss"
LOTE = 1000  # máximo de filas que devuelve la API por petición
NEGRITA = Font(bold=True)
TITULO = Font(bold=True, size=13)


class Api:
    def __init__(self, url: str, key: str):
        self.base = f"{url.rstrip('/')}/rest/v1"
        self.http = httpx.Client(headers={"apikey": key, "Authorization": f"Bearer {key}"}, timeout=60)

    def esquema(self) -> dict[str, dict[str, str]]:
        """{tabla: {columna: tipo}} de todas las tablas expuestas del esquema public."""
        r = self.http.get(self.base + "/", headers={"Accept": "application/openapi+json"})
        r.raise_for_status()
        return {
            tabla: {col: (p.get("format") or p.get("type") or "") for col, p in d.get("properties", {}).items()}
            for tabla, d in sorted(r.json().get("definitions", {}).items())
        }

    def leer_tabla(self, tabla: str, columnas: list[str]) -> list[dict]:
        orden = "id" if "id" in columnas else columnas[0]
        filas: list[dict] = []
        while True:
            r = self.http.get(f"{self.base}/{tabla}", params={
                "select": "*", "order": f"{orden}.asc", "offset": len(filas), "limit": LOTE,
            })
            r.raise_for_status()
            lote = r.json()
            filas.extend(lote)
            if len(lote) < LOTE:
                return filas


def a_madrid(valor: str | None) -> datetime | None:
    """'2026-10-22T21:49:03.12+00:00' -> datetime naive en hora de Madrid (Excel no admite zonas)."""
    if not valor:
        return None
    fecha = datetime.fromisoformat(valor.replace("Z", "+00:00")).astimezone(MADRID)
    return fecha.replace(tzinfo=None, microsecond=0)


def celda(valor, tipo: str):
    if valor is None:
        return None
    if tipo.startswith("timestamp"):
        return a_madrid(valor)
    if isinstance(valor, list):
        return ", ".join(str(v) for v in valor)
    if isinstance(valor, dict):
        return str(valor)
    if isinstance(valor, str):
        return ILLEGAL_CHARACTERS_RE.sub("", valor)
    return valor


def escribir_tabla(ws, fila_inicio: int, cabecera: list[str], filas: list[list]) -> int:
    """Escribe una tabla con cabecera en negrita. Devuelve la siguiente fila libre."""
    for c, nombre in enumerate(cabecera, 1):
        ws.cell(fila_inicio, c, nombre).font = NEGRITA
    for f, fila in enumerate(filas, fila_inicio + 1):
        for c, valor in enumerate(fila, 1):
            cel = ws.cell(f, c, valor)
            if isinstance(valor, datetime):
                cel.number_format = FORMATO_FECHA
    return fila_inicio + len(filas) + 2


def ajustar_anchos(ws, maximo: int = 60) -> None:
    anchos: dict[int, int] = {}
    for fila in ws.iter_rows():
        for cel in fila:
            if cel.value is not None:
                largo = 19 if isinstance(cel.value, datetime) else len(str(cel.value))
                anchos[cel.column] = max(anchos.get(cel.column, 0), largo)
    for col, ancho in anchos.items():
        ws.column_dimensions[get_column_letter(col)].width = min(max(ancho + 2, 8), maximo)


def hoja_resumen(ws, generado: datetime, datos: dict[str, list[dict]]) -> None:
    ws.cell(1, 1, "Snapshot de la base de datos").font = TITULO
    ws.cell(2, 1, "Generado (hora de Madrid)")
    ws.cell(2, 2, generado).number_format = FORMATO_FECHA
    fila = 4

    ws.cell(fila, 1, "Filas por tabla").font = TITULO
    fila = escribir_tabla(ws, fila + 1, ["Tabla", "Filas"], [[t, len(f)] for t, f in datos.items()])

    fuentes = {f["id"]: f["nombre"] for f in datos.get("fuentes", [])}

    ws.cell(fila, 1, "Artículos por fuente y día (día de publicación, hora de Madrid)").font = TITULO
    conteo = Counter(
        (a_madrid(a["publicado_en"]).date(), fuentes.get(a["fuente_id"], a["fuente_id"]))
        for a in datos.get("articulos", []) if a.get("publicado_en")
    )
    nombres = sorted({n for _, n in conteo}, key=str)
    dias = sorted({d for d, _ in conteo}, reverse=True)
    filas = [[d, *[conteo.get((d, n), 0) for n in nombres], sum(conteo.get((d, n), 0) for n in nombres)] for d in dias]
    inicio_dias = fila + 2
    fila = escribir_tabla(ws, fila + 1, ["Día", *nombres, "Total"], filas)
    for f in range(inicio_dias, inicio_dias + len(dias)):
        ws.cell(f, 1).number_format = "dd/mm/yyyy"

    ws.cell(fila, 1, "Últimas 10 ejecuciones").font = TITULO
    ultimas = sorted(datos.get("ejecuciones", []), key=lambda e: e["inicio"], reverse=True)[:10]
    escribir_tabla(
        ws, fila + 1,
        ["Inicio", "Fin", "Franja", "Fuente", "Estado", "Leídos", "Nuevos", "Páginas", "Mensaje"],
        [[a_madrid(e["inicio"]), a_madrid(e["fin"]), e["franja"], fuentes.get(e["fuente_id"], e["fuente_id"]),
          e["estado"], e["n_leidos"], e["n_nuevos"], e["paginas_leidas"], celda(e["error_msg"], "text")]
         for e in ultimas],
    )
    ajustar_anchos(ws, maximo=80)


def main() -> int:
    load_dotenv(RAIZ / ".env")
    url, key = os.environ.get("SUPABASE_URL"), os.environ.get("SUPABASE_SERVICE_ROLE_KEY")
    if not url or not key:
        print("Faltan SUPABASE_URL / SUPABASE_SERVICE_ROLE_KEY en .env", file=sys.stderr)
        return 1

    api = Api(url, key)
    esquema = api.esquema()
    datos = {}
    for tabla, columnas in esquema.items():
        datos[tabla] = api.leer_tabla(tabla, list(columnas))
        print(f"{tabla}: {len(datos[tabla])} filas")

    generado = datetime.now(MADRID).replace(tzinfo=None, microsecond=0)
    wb = Workbook()
    hoja_resumen(wb.active, generado, datos)
    wb.active.title = "resumen"

    for tabla, columnas in esquema.items():
        ws = wb.create_sheet(tabla[:31])  # límite de Excel para nombres de hoja
        nombres = list(columnas)
        escribir_tabla(ws, 1, nombres, [[celda(f.get(c), columnas[c]) for c in nombres] for f in datos[tabla]])
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        ajustar_anchos(ws)

    destino = RAIZ / "exports" / f"snapshot_{generado:%Y-%m-%d_%H%M}.xlsx"
    destino.parent.mkdir(exist_ok=True)
    wb.save(destino)
    print(f"Excel generado: {destino}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
